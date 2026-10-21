from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

def rupiah(nilai: int) -> str:
    return "Rp" + f"{nilai:,}".replace(",", ".")

class InvoicePdf:
    """PDF invoice sederhana dengan font bawaan fpdf2 (tanpa karakter non-Latin)."""

    def build(self, data: dict) -> bytes:
        pdf = FPDF(format="A4")
        pdf.set_auto_page_break(True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 18)
        pdf.cell(0, 10, "Cakra Space", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 6, data.get("building_name") or "Coworking Space Management", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 8, f"INVOICE {data['number']}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        for label, nilai in (
            ("Tanggal", data["issued_at"]),
            ("Status", data["status"].upper()),
            ("Kepada", f"{data['customer_name']} ({data['customer_email']})"),
        ):
            pdf.cell(35, 6, label)
            pdf.cell(0, 6, f": {nilai}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(130, 8, "Keterangan", border=1)
        pdf.cell(0, 8, "Jumlah", border=1, align="R", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        baris = [(data["description"], data["subtotal"])]
        if data["discount"]:
            baris.append(("Diskon", -data["discount"]))
        baris.append((f"Pajak ({data['tax_rate']}%)", data["tax"]))
        for ket, nilai in baris:
            teks = rupiah(abs(nilai))
            pdf.cell(130, 8, ket, border=1)
            pdf.cell(0, 8, f"-{teks}" if nilai < 0 else teks, border=1, align="R", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(130, 9, "Total", border=1)
        pdf.cell(0, 9, rupiah(data["total"]), border=1, align="R", new_x="LMARGIN", new_y="NEXT")
        return bytes(pdf.output())

class TableExporter:
    """Ekspor laporan tabular ke xlsx atau PDF."""

    def xlsx(self, title: str, headers: list[str], rows: list[list]) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def pdf(self, title: str, headers: list[str], rows: list[list]) -> bytes:
        pdf = FPDF(orientation="L", format="A4")
        pdf.set_auto_page_break(True, margin=12)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
        lebar = (pdf.w - 20) / max(len(headers), 1)
        pdf.set_font("Helvetica", "B", 9)
        for h in headers:
            pdf.cell(lebar, 7, str(h), border=1)
        pdf.ln()
        pdf.set_font("Helvetica", "", 8)
        for row in rows:
            for nilai in row:
                pdf.cell(lebar, 6, str(nilai)[:40], border=1)
            pdf.ln()
        return bytes(pdf.output())
